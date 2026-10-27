import os
import traceback
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from speedcopy import copyfile

import pandas as pd


def copy_spk_files(incoming_data: dict, current_progress, now_doc, all_doc, line_doing, line_progress, progress_value,
                   event, window_check, info_value) -> dict:
    logging = incoming_data['logging']
    errors = []
    too_many_sn = []
    not_find_files = []
    try:
        # finish_sn_list = os.listdir(incoming_data['set_folder'])
        df = pd.read_excel(Path(incoming_data['excel_file']), sheet_name=0, header=None, dtype=str)
        df[1] = df[1].astype(str)
        df[3] = df[3].astype(str)
        found_row = []
        for spk_file in Path(incoming_data['spk_folder']).rglob('*.txt'):
            spk = spk_file.stem.partition('_')[0].upper()
            if spk != 'SPK' and spk != 'СПК':
                logging.info(f"Файл {spk_file.name} не имеет в названии «SPK» или «СПК», пропускаем")
                continue
            if spk == 'SPK':
                spk = 'СПК'
            sn = spk_file.stem.partition('_')[2]
            if sn.startswith("00") is False:
                sn = "00" + sn
            # if sn in finish_sn_list:
            #     copy_file = Path(incoming_data['set_folder'], sn, spk + '_' + sn + spk_file.suffix)
            #     if copy_file.exists():
            #         logging.warning(f"Для файла {spk_file.name} в папке {sn} уже есть СПК, пропускаем")
            #     else:
            #         copyfile(str(spk_file), str(copy_file))
            #     continue
            for column in df.columns[3:]:
                try:
                    indexes = df.index[df[column] == sn].tolist()
                    if not indexes:
                        continue
                    if len(indexes) > 1:
                        err = f"У файла {spk_file.name} в столбце {get_column_letter(column)} слишком много совпадений"
                        logging.warning(err)
                        too_many_sn.append(err)
                        break
                    sn = df.iloc[indexes[0], 3]
                    copy_file = Path(incoming_data['set_folder'], sn, spk + '_' + sn + spk_file.suffix)
                    if copy_file.exists():
                        logging.warning(f"Для файла {spk_file.name} в папке {sn} уже есть СПК, пропускаем")
                        found_row.append(indexes[0] + 1)
                        break
                    copyfile(str(spk_file), str(copy_file))
                    found_row.append(int(indexes[0]) + 1)
                    break
                except BaseException as ex:
                    logging.error(f"Ошибка при копировании файла {spk_file.name} - {ex}")
                    logging.error(traceback.format_exc())
                    errors.append(f"Ошибка при копировании файла {spk_file.name}")
                    break
            else:
                logging.info(f"Файл {spk_file.name} не найден в excel")
                not_find_files.append(spk_file.name)
        found_row = set(found_row)
        xlsx_path = Path(incoming_data['excel_file'])
        wb = load_workbook(str(xlsx_path))
        ws = wb.active
        green = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
        red = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
        for row in range(3, ws.max_row + 1):
            cell = ws[f"D{row}"]
            cell.fill = green if row in found_row else red
        while True:
            try:
                wb.save(str(xlsx_path))
                break
            except PermissionError:
                info_value.emit('Вопрос?',
                                f"Файл {xlsx_path.name} открыт. Для сохранения ошибок в файл закройте "
                                f"его и нажмите «Да». Если нажать «Нет» ошибки не будут выделены!",
                                "")
                event.clear()
                event.wait()
            if window_check.answer is False:
                break
        os.startfile(xlsx_path)
        if errors or not_find_files or too_many_sn:
            file_name = Path(incoming_data['excel_file']).stem + '.txt'
            file_path = incoming_data['default_path']
            with open(Path(file_path, file_name), 'w', encoding='utf-8') as f:
                for item in errors + not_find_files + too_many_sn:
                    f.write(f'{item}\n')
            os.startfile(Path(file_path, file_name))
        return {'status': 'success', 'text': '', 'trace': ''}
    except BaseException as ex:
        return {'status': 'error', 'text': ex, 'trace': traceback.format_exc()}
